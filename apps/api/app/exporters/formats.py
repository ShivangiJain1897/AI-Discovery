"""Format renderers. Each takes a Document and returns bytes."""

from __future__ import annotations

import csv
import io
import json

from app.exporters.model import Document

_STATE_LABEL = {
    "known": "",
    "likely": "Likely — inferred from evidence",
    "hypothesis": "Hypothesis — not yet validated",
    "unknown": "Unknown — not established by this research",
}


def _state_note(section) -> str | None:
    parts = [_STATE_LABEL.get(section.knowledge_state, ""), section.note or ""]
    joined = " · ".join(p for p in parts if p)
    return joined or None


# ── Markdown ─────────────────────────────────────────────────


def to_markdown(doc: Document) -> bytes:
    out: list[str] = [f"# {doc.title}", ""]
    if doc.subtitle:
        out += [f"*{doc.subtitle}*", ""]
    if doc.summary:
        out += [doc.summary, ""]

    if doc.metadata:
        out += ["| Field | Value |", "| --- | --- |"]
        out += [f"| {k} | {v} |" for k, v in doc.metadata.items()]
        out.append("")

    for section in doc.sections:
        out += [f"## {section.heading}", ""]
        note = _state_note(section)
        if note:
            out += [f"> **{note}**", ""]
        if section.body:
            out += [section.body, ""]
        if section.points:
            out += [f"- {p}" for p in section.points] + [""]
        if section.rows and section.columns:
            out += [
                "| " + " | ".join(section.columns) + " |",
                "| " + " | ".join("---" for _ in section.columns) + " |",
            ]
            out += [
                "| "
                + " | ".join(
                    str(row.get(col, "")).replace("|", "\\|").replace("\n", " ")
                    for col in section.columns
                )
                + " |"
                for row in section.rows
            ]
            out.append("")
        if section.evidence_refs:
            out += [f"*Evidence: {', '.join(section.evidence_refs)}*", ""]

    if doc.open_questions:
        out += ["## Open Questions", ""] + [f"- {q}" for q in doc.open_questions] + [""]

    if doc.citations:
        out += ["## Research References", ""]
        for citation in doc.citations:
            link = f"[{citation.source_title}]({citation.source_url})" if citation.source_url else citation.source_title
            meta = " · ".join(
                filter(None, [citation.publisher, citation.publication_date, citation.tier])
            )
            out += [
                f"**[{citation.ref}]** {citation.statement}",
                "",
                f"> {citation.excerpt}",
                "",
                f"{link} — {meta} — {citation.evidence_type}, {citation.strength} evidence",
                "",
            ]

    out += ["---", f"*Generated {doc.generated_at:%Y-%m-%d %H:%M UTC}*"]
    return "\n".join(out).encode("utf-8")


# ── JSON ─────────────────────────────────────────────────────


def to_json(doc: Document) -> bytes:
    payload = {
        "title": doc.title,
        "subtitle": doc.subtitle,
        "summary": doc.summary,
        "metadata": doc.metadata,
        "generated_at": doc.generated_at.isoformat(),
        "sections": [
            {
                "heading": s.heading,
                "body": s.body,
                "points": s.points,
                "columns": s.columns,
                "rows": s.rows,
                "evidence_refs": s.evidence_refs,
                "knowledge_state": s.knowledge_state,
                "note": s.note,
            }
            for s in doc.sections
        ],
        "open_questions": doc.open_questions,
        "citations": [
            {
                "ref": c.ref,
                "statement": c.statement,
                "excerpt": c.excerpt,
                "source_title": c.source_title,
                "source_url": c.source_url,
                "publisher": c.publisher,
                "publication_date": c.publication_date,
                "tier": c.tier,
                "evidence_type": c.evidence_type,
                "strength": c.strength,
            }
            for c in doc.citations
        ],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


# ── CSV ──────────────────────────────────────────────────────


def to_csv(doc: Document) -> bytes:
    """Flatten every table in the document into one CSV."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    wrote = False

    for section in doc.sections:
        if not (section.rows and section.columns):
            continue
        if wrote:
            writer.writerow([])
        writer.writerow([section.heading])
        writer.writerow(section.columns)
        for row in section.rows:
            writer.writerow([row.get(col, "") for col in section.columns])
        wrote = True

    if not wrote:
        writer.writerow(["Heading", "Body", "Points", "Evidence"])
        for section in doc.sections:
            writer.writerow([
                section.heading, section.body,
                " | ".join(section.points), ", ".join(section.evidence_refs),
            ])

    return buffer.getvalue().encode("utf-8")


# ── XLSX ─────────────────────────────────────────────────────


def to_xlsx(doc: Document) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="0C5B4E")
    wrap = Alignment(vertical="top", wrap_text=True)

    def write_table(sheet, columns: list[str], rows: list[dict]) -> None:
        sheet.append(columns)
        for index in range(1, len(columns) + 1):
            cell = sheet.cell(row=1, column=index)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        for row in rows:
            sheet.append([str(row.get(col, "")) for col in columns])
        for index, column in enumerate(columns, start=1):
            widest = max([len(column)] + [len(str(r.get(column, ""))) for r in rows] or [10])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 12), 60)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        sheet.freeze_panes = "A2"

    # Overview sheet first: metadata, then the summary.
    overview = workbook.create_sheet("Overview")
    overview.append(["Field", "Value"])
    for index in (1, 2):
        overview.cell(row=1, column=index).font = header_font
        overview.cell(row=1, column=index).fill = header_fill
    overview.append(["Title", doc.title])
    overview.append(["Type", doc.subtitle])
    for key, value in doc.metadata.items():
        overview.append([key, value])
    overview.append(["Summary", doc.summary])
    overview.append(["Generated", doc.generated_at.strftime("%Y-%m-%d %H:%M UTC")])
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 100
    for row in overview.iter_rows(min_row=2):
        row[1].alignment = wrap

    used_names: set[str] = {"Overview"}
    for section in doc.sections:
        if not (section.rows and section.columns):
            continue
        # Excel sheet names: 31 chars, no []:*?/\ and must be unique.
        base = "".join(c for c in section.heading if c not in "[]:*?/\\")[:28] or "Section"
        name, suffix = base, 2
        while name in used_names:
            name = f"{base[:26]}-{suffix}"
            suffix += 1
        used_names.add(name)
        write_table(workbook.create_sheet(name), section.columns, section.rows)

    prose = [s for s in doc.sections if s.has_content and not s.rows]
    if prose:
        sheet = workbook.create_sheet("Narrative")
        write_table(
            sheet,
            ["Section", "Content", "Knowledge state", "Evidence"],
            [
                {
                    "Section": s.heading,
                    "Content": s.body or "\n".join(f"• {p}" for p in s.points),
                    "Knowledge state": s.knowledge_state,
                    "Evidence": ", ".join(s.evidence_refs),
                }
                for s in prose
            ],
        )

    if doc.citations:
        write_table(
            workbook.create_sheet("References"),
            ["Ref", "Statement", "Excerpt", "Source", "Publisher", "Published",
             "Tier", "Evidence type", "Strength", "URL"],
            [
                {
                    "Ref": c.ref, "Statement": c.statement, "Excerpt": c.excerpt,
                    "Source": c.source_title, "Publisher": c.publisher or "",
                    "Published": c.publication_date or "", "Tier": c.tier,
                    "Evidence type": c.evidence_type, "Strength": c.strength,
                    "URL": c.source_url or "",
                }
                for c in doc.citations
            ],
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── DOCX ─────────────────────────────────────────────────────


def to_docx(doc: Document) -> bytes:
    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = DocxDocument()

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    document.add_heading(doc.title, level=0)
    if doc.subtitle:
        subtitle = document.add_paragraph(doc.subtitle)
        subtitle.runs[0].italic = True
        subtitle.runs[0].font.color.rgb = RGBColor(0x5B, 0x5F, 0x66)

    if doc.metadata:
        table = document.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for key, value in doc.metadata.items():
            cells = table.add_row().cells
            cells[0].text = key
            cells[0].paragraphs[0].runs[0].bold = True
            cells[1].text = str(value)
        document.add_paragraph()

    if doc.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(doc.summary)

    for section in doc.sections:
        document.add_heading(section.heading, level=1)

        note = _state_note(section)
        if note:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(note)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)

        if section.body:
            document.add_paragraph(section.body)
        for point in section.points:
            document.add_paragraph(str(point), style="List Bullet")

        if section.rows and section.columns:
            table = document.add_table(rows=1, cols=len(section.columns))
            table.style = "Light Grid Accent 1"
            for index, column in enumerate(section.columns):
                cell = table.rows[0].cells[index]
                cell.text = column
                cell.paragraphs[0].runs[0].bold = True
            for row in section.rows:
                cells = table.add_row().cells
                for index, column in enumerate(section.columns):
                    cells[index].text = str(row.get(column, ""))
            document.add_paragraph()

        if section.evidence_refs:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"Evidence: {', '.join(section.evidence_refs)}")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x0C, 0x5B, 0x4E)

    if doc.open_questions:
        document.add_heading("Open Questions", level=1)
        for question in doc.open_questions:
            document.add_paragraph(str(question), style="List Bullet")

    if doc.citations:
        document.add_page_break()
        document.add_heading("Research References", level=1)
        for citation in doc.citations:
            paragraph = document.add_paragraph()
            ref_run = paragraph.add_run(f"[{citation.ref}] ")
            ref_run.bold = True
            ref_run.font.color.rgb = RGBColor(0x0C, 0x5B, 0x4E)
            paragraph.add_run(citation.statement)

            excerpt = document.add_paragraph(f"\u201c{citation.excerpt}\u201d")
            excerpt.paragraph_format.left_indent = Pt(24)
            excerpt.runs[0].italic = True
            excerpt.runs[0].font.size = Pt(9)

            meta_parts = filter(
                None,
                [
                    citation.source_title, citation.publisher, citation.publication_date,
                    citation.tier, f"{citation.evidence_type}, {citation.strength} evidence",
                    citation.source_url,
                ],
            )
            meta = document.add_paragraph(" · ".join(meta_parts))
            meta.paragraph_format.left_indent = Pt(24)
            meta.runs[0].font.size = Pt(8)
            meta.runs[0].font.color.rgb = RGBColor(0x5B, 0x5F, 0x66)

    footer = document.add_paragraph(f"Generated {doc.generated_at:%Y-%m-%d %H:%M UTC}")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x8A, 0x8F, 0x97)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── PDF ──────────────────────────────────────────────────────


def to_pdf(doc: Document) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    def escape(text) -> str:
        return (
            str(text)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    pine = colors.HexColor("#0C5B4E")
    muted = colors.HexColor("#5B5F66")
    amber = colors.HexColor("#8A5A00")

    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=22, leading=26, textColor=colors.HexColor("#17181B")),
        "subtitle": ParagraphStyle("st", parent=base["Normal"], fontSize=11, textColor=muted, spaceAfter=14),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontSize=14, leading=18, textColor=pine, spaceBefore=16, spaceAfter=6),
        "body": ParagraphStyle("b", parent=base["BodyText"], fontSize=10, leading=14),
        "bullet": ParagraphStyle("bu", parent=base["BodyText"], fontSize=10, leading=14, leftIndent=14, bulletIndent=4),
        "note": ParagraphStyle("n", parent=base["Normal"], fontSize=8.5, textColor=amber, spaceAfter=4),
        "ref": ParagraphStyle("r", parent=base["Normal"], fontSize=8.5, textColor=pine, spaceBefore=4),
        "cite": ParagraphStyle("c", parent=base["Normal"], fontSize=8.5, textColor=muted, leftIndent=16),
        "excerpt": ParagraphStyle("e", parent=base["Normal"], fontSize=9, leading=12, leftIndent=16, textColor=colors.HexColor("#17181B")),
        "footer": ParagraphStyle("f", parent=base["Normal"], fontSize=8, textColor=muted, alignment=TA_CENTER),
    }

    story = [Paragraph(escape(doc.title), styles["title"])]
    if doc.subtitle:
        story.append(Paragraph(escape(doc.subtitle), styles["subtitle"]))

    def make_table(columns: list[str], rows: list[list[str]], widths=None) -> Table:
        data = [[Paragraph(f"<b>{escape(c)}</b>", styles["cite"]) for c in columns]]
        data += [[Paragraph(escape(cell), styles["cite"]) for cell in row] for row in rows]
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), pine),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CFCBC3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F2EE")]),
            ])
        )
        return table

    if doc.metadata:
        width = 6.5 * inch
        story.append(
            make_table(
                ["Field", "Value"],
                [[k, str(v)] for k, v in doc.metadata.items()],
                widths=[1.8 * inch, width - 1.8 * inch],
            )
        )
        story.append(Spacer(1, 12))

    if doc.summary:
        story += [Paragraph("Summary", styles["h1"]), Paragraph(escape(doc.summary), styles["body"])]

    for section in doc.sections:
        story.append(Paragraph(escape(section.heading), styles["h1"]))
        note = _state_note(section)
        if note:
            story.append(Paragraph(escape(note), styles["note"]))
        if section.body:
            story.append(Paragraph(escape(section.body), styles["body"]))
        for point in section.points:
            story.append(Paragraph(escape(point), styles["bullet"], bulletText="•"))
        if section.rows and section.columns:
            available = 6.5 * inch
            widths = [available / len(section.columns)] * len(section.columns)
            story.append(Spacer(1, 6))
            story.append(
                make_table(
                    section.columns,
                    [[str(r.get(c, "")) for c in section.columns] for r in section.rows],
                    widths=widths,
                )
            )
            story.append(Spacer(1, 6))
        if section.evidence_refs:
            story.append(
                Paragraph(f"Evidence: {escape(', '.join(section.evidence_refs))}", styles["ref"])
            )

    if doc.open_questions:
        story.append(Paragraph("Open Questions", styles["h1"]))
        for question in doc.open_questions:
            story.append(Paragraph(escape(question), styles["bullet"], bulletText="•"))

    if doc.citations:
        story.append(PageBreak())
        story.append(Paragraph("Research References", styles["h1"]))
        for citation in doc.citations:
            story.append(
                Paragraph(
                    f'<font color="#0C5B4E"><b>[{escape(citation.ref)}]</b></font> '
                    f"{escape(citation.statement)}",
                    styles["body"],
                )
            )
            story.append(Paragraph(f"\u201c{escape(citation.excerpt)}\u201d", styles["excerpt"]))
            meta = " · ".join(
                filter(None, [
                    citation.source_title, citation.publisher, citation.publication_date,
                    citation.tier, f"{citation.evidence_type}, {citation.strength} evidence",
                    citation.source_url,
                ])
            )
            story.append(Paragraph(escape(meta), styles["cite"]))
            story.append(Spacer(1, 8))

    story.append(Spacer(1, 18))
    story.append(
        Paragraph(f"Generated {doc.generated_at:%Y-%m-%d %H:%M UTC}", styles["footer"])
    )

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.85 * inch, bottomMargin=0.85 * inch,
        title=doc.title, author="AI Product Discovery Platform",
    ).build(story)
    return buffer.getvalue()
