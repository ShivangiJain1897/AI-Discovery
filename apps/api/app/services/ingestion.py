"""Internal document ingestion.

Uploaded material — interview transcripts, support exports, existing PRDs — is
parsed here and becomes evidence labelled INTERNAL. That label is load-bearing:
internal evidence is never used to build a public search query, and it is
filterable separately in the evidence library so a reader can always tell
whether a finding rests on public research or on the organization's own data.
"""

from __future__ import annotations

import csv
import io
import json
import logging

from sqlalchemy.orm import Session

from app.config import settings
from app.db.models import Evidence, Project, Source, UploadedDocument
from app.domain.contracts import EvidenceExtraction
from app.domain.enums import EvidenceOrigin, EvidenceStrength, ResearchType, SourceTier, SourceType
from app.orchestration.agents import run_stage
from app.orchestration.serializers import context_dict, source_dict
from app.services import refs, safety, storage

logger = logging.getLogger(__name__)

SUPPORTED = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}

#: What kind of internal material this is, which sets the evidence's research
#: type — an interview transcript is user research, a ticket export is VOC.
DOCUMENT_KINDS = {
    "interview_transcript": ResearchType.USER,
    "meeting_notes": ResearchType.WORKFLOW_OPERATIONAL,
    "survey_results": ResearchType.USER,
    "support_tickets": ResearchType.VOICE_OF_CUSTOMER,
    "customer_feedback": ResearchType.VOICE_OF_CUSTOMER,
    "existing_prd": ResearchType.PROBLEM_DOMAIN,
    "strategy_document": ResearchType.BUSINESS_COMMERCIAL,
    "research_report": ResearchType.PROBLEM_DOMAIN,
    "other": ResearchType.PROBLEM_DOMAIN,
}


class UnsupportedDocument(Exception):
    pass


def extension_of(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def parse(filename: str, data: bytes) -> tuple[str, int]:
    """Extract text from an uploaded file. Returns (text, page_count)."""
    extension = extension_of(filename)

    if extension == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages), len(pages)

    if extension == "docx":
        from docx import Document as DocxDocument

        document = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts), 0

    if extension == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts), len(workbook.worksheets)

    if extension == "csv":
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        return "\n".join(" | ".join(row) for row in reader), 0

    if extension == "json":
        try:
            parsed = json.loads(data.decode("utf-8", errors="replace"))
        except json.JSONDecodeError as exc:
            raise UnsupportedDocument(f"Invalid JSON: {exc}") from exc
        return json.dumps(parsed, indent=2)[:400000], 0

    if extension in ("txt", "md", "markdown"):
        return data.decode("utf-8", errors="replace"), 0

    raise UnsupportedDocument(
        f"Cannot parse '.{extension}'. Supported: {', '.join(sorted(SUPPORTED))}"
    )


def ingest(
    session: Session,
    project: Project,
    *,
    filename: str,
    data: bytes,
    document_kind: str = "other",
) -> UploadedDocument:
    """Store an uploaded document, parse it and scan it for identifiers."""
    text, pages = parse(filename, data)
    scan = safety.scan(text)

    ref = refs.allocate_one(session, project.id, "document")
    key = f"projects/{project.id}/uploads/{ref}-{filename}"
    storage.get_storage().write(key, data)

    document = UploadedDocument(
        project_id=project.id,
        ref=ref,
        filename=filename[:400],
        content_type=SUPPORTED.get(extension_of(filename), "application/octet-stream"),
        size_bytes=len(data),
        storage_key=key,
        document_kind=document_kind,
        origin=EvidenceOrigin.INTERNAL,
        extracted_text=text[:400000],
        page_count=pages,
        status="parsed",
        phi_scan=scan.as_dict(),
    )
    session.add(document)
    session.flush()

    if scan.found:
        logger.warning(
            "Uploaded document %s contains identifiers: %s", ref, ", ".join(scan.kinds)
        )
    return document


def extract_evidence(
    session: Session, project: Project, document: UploadedDocument
) -> list[Evidence]:
    """Turn a parsed internal document into INTERNAL evidence."""
    if not document.extracted_text:
        return []

    research_type = DOCUMENT_KINDS.get(document.document_kind, ResearchType.PROBLEM_DOMAIN)

    source = Source(
        project_id=project.id,
        document_id=document.id,
        ref=refs.allocate_one(session, project.id, "source"),
        title=document.filename,
        url=None,
        publisher="Internal",
        source_type=_source_type_for(document.document_kind),
        # Internal first-hand material is primary evidence about this
        # organization, even though it is not a public authority.
        tier=SourceTier.TIER_1_AUTHORITATIVE,
        origin=EvidenceOrigin.INTERNAL,
        research_types=[research_type.value],
        snippet=document.extracted_text[:4000],
        raw_content=document.extracted_text[:60000],
    )
    session.add(source)
    session.flush()

    extraction, _ = run_stage(
        "evidence_extractor",
        {
            "source": source_dict(source),
            "research_question": project.question,
            "research_type": research_type.value,
            "normalized_context": context_dict(project.context),
        },
        schema=EvidenceExtraction,
    )

    source.assessment = extraction.source_assessment or None
    items = extraction.items[: settings.max_evidence_per_source]
    if not items:
        document.status = "no_evidence"
        session.flush()
        return []

    created: list[Evidence] = []
    evidence_refs = refs.allocate(session, project.id, "evidence", len(items))
    for item, ref in zip(items, evidence_refs, strict=True):
        evidence = Evidence(
            project_id=project.id,
            source_id=source.id,
            ref=ref,
            research_type=research_type,
            statement=item.statement,
            excerpt=item.excerpt,
            evidence_type=item.evidence_type,
            origin=EvidenceOrigin.INTERNAL,
            population=item.population,
            geography=item.geography,
            period=item.period,
            quantities=item.quantities,
            tags=item.theme_candidates,
            relevance=item.relevance,
            limitations=item.limitations,
            strength=EvidenceStrength.MODERATE,
            strength_reasoning=(
                "Internal first-hand material: authoritative about this organization, "
                "not generalizable beyond it."
            ),
        )
        session.add(evidence)
        created.append(evidence)

    document.status = "ingested"
    session.flush()
    return created


def _source_type_for(kind: str) -> SourceType:
    return {
        "interview_transcript": SourceType.INTERVIEW_TRANSCRIPT,
        "support_tickets": SourceType.SUPPORT_TICKETS,
        "customer_feedback": SourceType.SUPPORT_TICKETS,
        "survey_results": SourceType.SURVEY,
    }.get(kind, SourceType.INTERNAL_DOCUMENT)
