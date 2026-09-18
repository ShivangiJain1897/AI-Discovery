"""Prompt library integrity and native document generation."""

from __future__ import annotations

import io

import pytest

from app.domain.contracts import Contract
from app.exporters.model import DocCitation, DocSection, Document
from app.exporters.registry import EXPORTERS, export_artifact
from app.prompts import get_prompt_library
from app.services import artifact_templates


class TestPromptLibrary:
    def test_all_stages_are_present(self):
        library = get_prompt_library()
        assert len(library.all()) == 31

    def test_every_prompt_declares_a_resolvable_schema(self):
        """A prompt naming a schema that no longer exists must fail at load,
        not mid-pipeline."""
        for template in get_prompt_library().all():
            schema = template.schema_class
            assert schema is not None, f"{template.key} declares no schema"
            assert issubclass(schema, Contract)

    def test_every_prompt_has_guardrails(self):
        for template in get_prompt_library().all():
            assert template.guardrails, f"{template.key} has no guardrails"

    def test_universal_preamble_is_prepended(self):
        system, _, _ = get_prompt_library().render("evidence_extractor", {})
        assert "EVIDENCE DISCIPLINE" in system
        assert "Never widen a population" in system

    def test_stage_instructions_follow_the_preamble(self):
        system, _, _ = get_prompt_library().render("swot_analysis", {})
        assert system.index("EVIDENCE DISCIPLINE") < system.index("STAGE: SWOT Analysis")

    def test_inputs_are_serialized_as_labelled_json(self):
        _, user, _ = get_prompt_library().render(
            "research_planner", {"normalized_context": {"domain": "Healthcare"}}
        )
        assert "## NORMALIZED CONTEXT" in user
        assert '"domain": "Healthcare"' in user

    def test_missing_declared_inputs_are_reported_not_hidden(self):
        _, user, _ = get_prompt_library().render("research_planner", {})
        assert "NOT SUPPLIED" in user

    def test_unknown_prompt_key_raises(self):
        with pytest.raises(KeyError):
            get_prompt_library().get("no_such_prompt")

    def test_prompts_are_checksummed_for_provenance(self):
        for template in get_prompt_library().all():
            assert len(template.checksum) == 64


class TestArtifactTemplates:
    def test_prd_has_all_twenty_three_sections(self):
        assert len(artifact_templates.section_headings("prd")) == 23

    def test_prd_opens_with_summary_and_closes_with_references(self):
        sections = artifact_templates.section_headings("prd")
        assert sections[0] == "Executive Summary"
        assert "Appendix / Research References" in sections[-1]

    def test_tabular_artifacts_offer_spreadsheet_exports(self):
        assert "xlsx" in artifact_templates.export_formats("use_case_catalog")

    def test_unknown_type_falls_back_to_prd_shape(self):
        assert artifact_templates.section_headings("not_a_type") == artifact_templates.PRD


@pytest.fixture
def document() -> Document:
    return Document(
        title="Cost Transparency PRD",
        subtitle="Product Requirements Document",
        summary="Members cannot estimate cost before care.",
        sections=[
            DocSection(
                heading="Problem Statement",
                body="Members lack cost estimates [E-001].",
                evidence_refs=["E-001"],
                knowledge_state="known",
            ),
            DocSection(
                heading="Use Cases",
                columns=["Use case", "Persona", "Evidence"],
                rows=[{"Use case": "Check cost", "Persona": "Member", "Evidence": "E-001"}],
            ),
            DocSection(heading="Success Metrics", knowledge_state="unknown", note="Unknown"),
        ],
        citations=[
            DocCitation(
                ref="E-001",
                statement="Cost estimates are unavailable before a visit.",
                excerpt="Beneficiaries report no way to determine cost in advance.",
                source_title="CMS Price Transparency",
                source_url="https://www.cms.gov/x",
                publisher="CMS",
                publication_date="2024-03-01",
                tier="Primary / Authoritative",
                strength="strong",
                evidence_type="fact",
            )
        ],
        open_questions=["Willingness to pay not established."],
        metadata={"Project": "Medicaid cost transparency", "Version": "v1"},
    )


class TestExports:
    @pytest.mark.parametrize("fmt", sorted(EXPORTERS))
    def test_every_format_produces_bytes(self, document, fmt):
        data, filename, media_type = export_artifact(document, fmt)
        assert len(data) > 0
        assert filename.startswith("cost-transparency-prd")
        assert media_type

    def test_docx_is_a_real_word_file(self, document):
        """Not an HTML screenshot — a real OOXML package."""
        import zipfile

        data, _, _ = export_artifact(document, "docx")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert "word/document.xml" in archive.namelist()

    def test_xlsx_is_a_real_workbook_with_a_references_sheet(self, document):
        from openpyxl import load_workbook

        data, _, _ = export_artifact(document, "xlsx")
        workbook = load_workbook(io.BytesIO(data))
        assert "Overview" in workbook.sheetnames
        assert "References" in workbook.sheetnames

    def test_pdf_has_a_pdf_header(self, document):
        data, _, _ = export_artifact(document, "pdf")
        assert data.startswith(b"%PDF")

    def test_citations_survive_markdown_export(self, document):
        data, _, _ = export_artifact(document, "markdown")
        text = data.decode()
        assert "[E-001]" in text
        assert "https://www.cms.gov/x" in text
        assert "Beneficiaries report no way" in text

    def test_knowledge_state_survives_export(self, document):
        """A reader of the exported file must still see which sections the
        research did not establish."""
        data, _, _ = export_artifact(document, "markdown")
        assert "Unknown" in data.decode()

    def test_tables_survive_csv_export(self, document):
        data, _, _ = export_artifact(document, "csv")
        text = data.decode()
        assert "Use case" in text and "Check cost" in text

    def test_json_export_round_trips(self, document):
        import json

        data, _, _ = export_artifact(document, "json")
        payload = json.loads(data)
        assert payload["title"] == "Cost Transparency PRD"
        assert payload["citations"][0]["ref"] == "E-001"
        assert len(payload["sections"]) == 3

    def test_unsupported_format_is_rejected_clearly(self, document):
        with pytest.raises(ValueError, match="Unsupported export format"):
            export_artifact(document, "pages")
